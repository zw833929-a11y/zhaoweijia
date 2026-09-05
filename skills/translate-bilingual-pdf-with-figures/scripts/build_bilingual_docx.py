#!/usr/bin/env python3
"""Build a polished paragraph-paired bilingual DOCX from the canonical JSON model."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = "1F5A73"
SOURCE_COLOR = RGBColor(35, 48, 56)
TRANSLATION_COLOR = RGBColor(31, 90, 115)
MUTED_COLOR = RGBColor(88, 100, 108)
SUPERSCRIPT_RE = re.compile(r"\^(?:\{([^}]+)\}|([+\-−]?\d+|[A-Za-z]+))")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a bilingual DOCX from a JSON model.")
    parser.add_argument("--model", type=Path, required=True, help="Canonical bilingual JSON model")
    parser.add_argument("--output", type=Path, required=True, help="Output DOCX")
    parser.add_argument("--latin-font", default="Aptos", help="Latin font name")
    parser.add_argument("--cjk-font", default="Noto Sans CJK SC", help="CJK font name")
    return parser.parse_args()


def set_run_font(run: Any, latin: str, cjk: str, size: float | None = None) -> None:
    run.font.name = latin
    if size is not None:
        run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), cjk)
    r_fonts.set(qn("w:cs"), latin)


def set_style_fonts(style: Any, latin: str, cjk: str, size: float | None = None) -> None:
    style.font.name = latin
    if size is not None:
        style.font.size = Pt(size)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), cjk)
    r_fonts.set(qn("w:cs"), latin)


def add_or_get_style(doc: Document, name: str, base: str = "Normal") -> Any:
    if name in doc.styles:
        return doc.styles[name]
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base]
    return style


def configure_styles(doc: Document, latin: str, cjk: str) -> None:
    normal = doc.styles["Normal"]
    set_style_fonts(normal, latin, cjk, 10.5)
    normal.font.color.rgb = SOURCE_COLOR
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)

    source = add_or_get_style(doc, "Source Paragraph")
    set_style_fonts(source, latin, cjk, 9.8)
    source.font.color.rgb = SOURCE_COLOR
    source.paragraph_format.line_spacing = 1.18
    source.paragraph_format.space_after = Pt(3)
    source.paragraph_format.keep_together = True

    translation = add_or_get_style(doc, "Translation Paragraph")
    set_style_fonts(translation, latin, cjk, 10.5)
    translation.font.color.rgb = TRANSLATION_COLOR
    translation.paragraph_format.line_spacing = 1.35
    translation.paragraph_format.space_after = Pt(10)
    translation.paragraph_format.keep_together = True

    for level in range(1, 4):
        source_heading = add_or_get_style(doc, f"Source Heading {level}", f"Heading {level}")
        set_style_fonts(source_heading, latin, cjk, {1: 16, 2: 13, 3: 11.5}[level])
        source_heading.font.color.rgb = RGBColor.from_string(ACCENT)
        source_heading.font.bold = True
        source_heading.paragraph_format.space_before = Pt({1: 14, 2: 11, 3: 8}[level])
        source_heading.paragraph_format.space_after = Pt(2)
        source_heading.paragraph_format.keep_with_next = True

        translation_heading = add_or_get_style(doc, f"Translation Heading {level}")
        set_style_fonts(translation_heading, latin, cjk, {1: 13.5, 2: 11.5, 3: 10.5}[level])
        translation_heading.font.color.rgb = TRANSLATION_COLOR
        translation_heading.font.bold = True
        translation_heading.paragraph_format.space_after = Pt(7)
        translation_heading.paragraph_format.keep_with_next = True

    source_caption = add_or_get_style(doc, "Source Caption")
    set_style_fonts(source_caption, latin, cjk, 8.5)
    source_caption.font.color.rgb = SOURCE_COLOR
    source_caption.font.italic = True
    source_caption.paragraph_format.line_spacing = 1.05
    source_caption.paragraph_format.space_after = Pt(2)
    source_caption.paragraph_format.keep_together = True

    translation_caption = add_or_get_style(doc, "Translation Caption")
    set_style_fonts(translation_caption, latin, cjk, 9)
    translation_caption.font.color.rgb = TRANSLATION_COLOR
    translation_caption.paragraph_format.line_spacing = 1.1
    translation_caption.paragraph_format.space_after = Pt(10)
    translation_caption.paragraph_format.keep_together = True

    reference = add_or_get_style(doc, "Reference")
    set_style_fonts(reference, latin, cjk, 8.5)
    reference.font.color.rgb = SOURCE_COLOR
    reference.paragraph_format.line_spacing = 1.0
    reference.paragraph_format.space_after = Pt(2)
    reference.paragraph_format.left_indent = Inches(0.18)
    reference.paragraph_format.first_line_indent = Inches(-0.18)

    note = add_or_get_style(doc, "Translator Note")
    set_style_fonts(note, latin, cjk, 9)
    note.font.color.rgb = MUTED_COLOR
    note.font.italic = True
    note.paragraph_format.space_before = Pt(6)
    note.paragraph_format.space_after = Pt(8)


def shade_paragraph(paragraph: Any, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_rich_text(paragraph: Any, text_value: str, latin: str, cjk: str) -> None:
    position = 0
    for match in SUPERSCRIPT_RE.finditer(text_value):
        if match.start() > position:
            run = paragraph.add_run(text_value[position:match.start()])
            set_run_font(run, latin, cjk)
        superscript = match.group(1) or match.group(2) or ""
        run = paragraph.add_run(superscript)
        set_run_font(run, latin, cjk)
        run.font.superscript = True
        position = match.end()
    if position < len(text_value):
        run = paragraph.add_run(text_value[position:])
        set_run_font(run, latin, cjk)


def bookmark_name(item_id: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_]", "_", item_id)[:24]
    digest = hashlib.sha1(item_id.encode("utf-8")).hexdigest()[:8]
    return f"item_{cleaned}_{digest}"


def add_bookmark(paragraph: Any, item_id: str, numeric_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(numeric_id))
    start.set(qn("w:name"), bookmark_name(item_id))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(numeric_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_field(paragraph: Any, field_code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {field_code} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, cached, end])


def add_page_number_footer(section: Any, latin: str, cjk: str) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = paragraph.add_run("Page ")
    set_run_font(prefix, latin, cjk, 8)
    add_field(paragraph, "PAGE")
    middle = paragraph.add_run(" / ")
    set_run_font(middle, latin, cjk, 8)
    add_field(paragraph, "NUMPAGES")
    for run in paragraph.runs:
        run.font.color.rgb = MUTED_COLOR


def add_header(section: Any, title: str, latin: str, cjk: str) -> None:
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(title)
    set_run_font(run, latin, cjk, 7.5)
    run.font.color.rgb = MUTED_COLOR


def validate_model(model: dict[str, Any], model_path: Path) -> None:
    if not isinstance(model.get("items"), list):
        raise ValueError("Model must contain an 'items' array")
    ids: set[str] = set()
    supported = {"heading", "paragraph", "figure", "equation"}
    for index, item in enumerate(model["items"]):
        if not isinstance(item, dict):
            raise ValueError(f"Item {index} is not an object")
        item_id = str(item.get("id", "")).strip()
        item_type = str(item.get("type", "")).strip()
        if not item_id:
            raise ValueError(f"Item {index} has no id")
        if item_id in ids:
            raise ValueError(f"Duplicate item id: {item_id}")
        ids.add(item_id)
        if item_type not in supported:
            raise ValueError(f"Unsupported item type for {item_id}: {item_type}")
        if item_type == "figure":
            for field in ("image", "caption_en", "caption_zh"):
                if not str(item.get(field, "")).strip():
                    raise ValueError(f"Figure {item_id} has no {field}")
            image_path = Path(str(item["image"]))
            if not image_path.is_absolute():
                image_path = model_path.parent / image_path
            if not image_path.is_file():
                raise ValueError(f"Figure image not found for {item_id}: {image_path}")
        else:
            for field in ("en", "zh"):
                if not str(item.get(field, "")).strip():
                    raise ValueError(f"Item {item_id} has no {field}")


def image_stream_and_ratio(image_path: Path) -> tuple[io.BytesIO, float]:
    with Image.open(image_path) as raw:
        image = ImageOps.exif_transpose(raw)
        if image.mode not in ("RGB", "RGBA", "L"):
            image = image.convert("RGB")
        width, height = image.size
        if width <= 0 or height <= 0:
            raise ValueError(f"Invalid image dimensions: {image_path}")
        stream = io.BytesIO()
        image.save(stream, format="PNG", optimize=True)
        stream.seek(0)
        return stream, width / height


def figure_width(aspect_ratio: float, max_width: float = 6.55, max_height: float = 7.75) -> Inches:
    width = max_width
    if width / aspect_ratio > max_height:
        width = max_height * aspect_ratio
    return Inches(max(1.2, min(max_width, width)))


def add_cover(doc: Document, metadata: dict[str, Any], latin: str, cjk: str) -> bool:
    title_en = str(metadata.get("title_en", "")).strip()
    title_zh = str(metadata.get("title_zh", "")).strip()
    if not title_en and not title_zh:
        return False

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(26)

    if title_en:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(10)
        run = paragraph.add_run(title_en)
        set_run_font(run, latin, cjk, 20)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(ACCENT)
    if title_zh:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(24)
        run = paragraph.add_run(title_zh)
        set_run_font(run, latin, cjk, 17)
        run.bold = True
        run.font.color.rgb = TRANSLATION_COLOR

    for key, label in (
        ("authors", "Authors / 作者"),
        ("journal", "Journal / 期刊"),
        ("doi", "DOI"),
        ("source_file", "Source / 来源文件"),
    ):
        value = str(metadata.get(key, "")).strip()
        if not value:
            continue
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, latin, cjk, 9)
        label_run.bold = True
        label_run.font.color.rgb = MUTED_COLOR
        value_run = paragraph.add_run(value)
        set_run_font(value_run, latin, cjk, 9)
        value_run.font.color.rgb = SOURCE_COLOR

    note_text = str(metadata.get("translator_note", "")).strip()
    if note_text:
        note = doc.add_paragraph(style="Translator Note")
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich_text(note, f"译者说明：{note_text}", latin, cjk)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(18)
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), ACCENT)
    borders.append(bottom)
    p_pr.append(borders)

    doc.add_page_break()
    return True


def add_pair(
    doc: Document,
    item: dict[str, Any],
    bookmark_id: int,
    latin: str,
    cjk: str,
) -> None:
    item_type = str(item["type"])
    if item_type == "heading":
        level = max(1, min(3, int(item.get("level", 1))))
        source = doc.add_paragraph(style=f"Source Heading {level}")
        add_rich_text(source, str(item["en"]), latin, cjk)
        add_bookmark(source, str(item["id"]), bookmark_id)
        translation = doc.add_paragraph(style=f"Translation Heading {level}")
        add_rich_text(translation, str(item["zh"]), latin, cjk)
        return

    source = doc.add_paragraph(style="Source Paragraph")
    add_rich_text(source, str(item["en"]), latin, cjk)
    add_bookmark(source, str(item["id"]), bookmark_id)
    translation = doc.add_paragraph(style="Translation Paragraph")
    add_rich_text(translation, str(item["zh"]), latin, cjk)


def add_figure(
    doc: Document,
    item: dict[str, Any],
    model_dir: Path,
    bookmark_id: int,
    latin: str,
    cjk: str,
) -> None:
    if bool(item.get("page_break")):
        doc.add_page_break()
    image_path = Path(str(item["image"]))
    if not image_path.is_absolute():
        image_path = model_dir / image_path
    stream, ratio = image_stream_and_ratio(image_path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(stream, width=figure_width(ratio))
    add_bookmark(paragraph, str(item["id"]), bookmark_id)

    source_caption = doc.add_paragraph(style="Source Caption")
    add_rich_text(source_caption, str(item["caption_en"]), latin, cjk)
    translation_caption = doc.add_paragraph(style="Translation Caption")
    add_rich_text(translation_caption, str(item["caption_zh"]), latin, cjk)


def iter_reference_strings(references: Iterable[Any]) -> Iterable[str]:
    for value in references:
        if isinstance(value, dict):
            text_value = str(value.get("text", "")).strip()
        else:
            text_value = str(value).strip()
        if text_value:
            yield text_value


def build_document(model: dict[str, Any], model_path: Path, latin: str, cjk: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    # Base geometry follows the compact_reference_guide preset. Typography below
    # is a named "academic bilingual" override for paired source/translation text.
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc, latin, cjk)
    metadata = model.get("metadata") if isinstance(model.get("metadata"), dict) else {}
    title_for_header = str(metadata.get("title_zh") or metadata.get("title_en") or "Bilingual article")
    add_header(section, title_for_header[:90], latin, cjk)
    add_page_number_footer(section, latin, cjk)

    properties = doc.core_properties
    properties.title = str(metadata.get("title_en") or metadata.get("title_zh") or "Bilingual article")
    properties.subject = "Paragraph-by-paragraph bilingual academic translation"
    properties.comments = "Generated from a verified structured article model."

    add_cover(doc, metadata, latin, cjk)

    bookmark_id = 1
    for item in model["items"]:
        if item["type"] == "figure":
            add_figure(doc, item, model_path.parent, bookmark_id, latin, cjk)
        else:
            add_pair(doc, item, bookmark_id, latin, cjk)
        bookmark_id += 1

    references = list(iter_reference_strings(model.get("references", [])))
    if references:
        heading_source = doc.add_paragraph(style="Source Heading 1")
        add_rich_text(heading_source, "References", latin, cjk)
        heading_translation = doc.add_paragraph(style="Translation Heading 1")
        add_rich_text(heading_translation, "参考文献（原文）", latin, cjk)
        for reference_text in references:
            paragraph = doc.add_paragraph(style="Reference")
            add_rich_text(paragraph, reference_text, latin, cjk)

    return doc


def main() -> int:
    args = parse_args()
    model_path = args.model.expanduser().resolve()
    output_path = args.output.expanduser().resolve()
    if not model_path.is_file():
        raise SystemExit(f"Model not found: {model_path}")
    model = json.loads(model_path.read_text(encoding="utf-8"))
    validate_model(model, model_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = build_document(model, model_path, args.latin_font, args.cjk_font)
    doc.save(output_path)
    figures = sum(1 for item in model["items"] if item["type"] == "figure")
    translations = sum(1 for item in model["items"] if item["type"] != "figure") + figures
    print(json.dumps({
        "output": str(output_path),
        "items": len(model["items"]),
        "translations": translations,
        "figures": figures,
        "references": len(list(iter_reference_strings(model.get("references", [])))),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
